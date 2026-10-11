"""
The bug fixes and agency features added in the October review: mail-merge drafting, placeholder
fallbacks, sending hours, delivery failures, replies read-or-unread, bounces, follow-ups (bulk
and by hand), the one-click pipeline, schedules, imports, exports and assorted hardening.
Everything external (AI, SMTP, IMAP) is mocked.
"""
import io
import json
import os
import sys
import unittest
from datetime import datetime, timezone, timedelta
from email.message import EmailMessage
from unittest import mock

sys.path.insert(0, os.path.dirname(__file__))
from common import Base, ApiBase, FakeCtx  # noqa: E402

import db  # noqa: E402
import enrichment  # noqa: E402
import exports  # noqa: E402
import filtering  # noqa: E402
import importer  # noqa: E402
import jobs  # noqa: E402
import llm  # noqa: E402
import outreach  # noqa: E402
import pipeline  # noqa: E402
import schema  # noqa: E402

OLD = "2026-01-01T10:00:00+00:00"


def emailed(test, name="Acme Plumbing", email="owner@acme.com", campaign="c", msgid="<m1@me.org>"):
    """A lead that was emailed long enough ago for a follow-up to be due."""
    lead_id = test.add(business_name=name, email=email, campaign=campaign)
    db.update_lead(lead_id, {"message": "Subject: Hello\n\nHi there", "send_status": "sent", "stage": "contacted",
                             "sent_at": OLD, "last_sent_at": OLD, "email_message_id": msgid})
    db.add_event(lead_id, "email_sent", json.dumps({"subject": "Hello", "message_id": msgid}))
    return lead_id


def reply_bytes(frm, in_reply_to="<m1@me.org>", body="Sounds good, call me.", msgid="<r1@them.com>", **headers):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"], m["Message-ID"] = frm, "me@example.org", "Re: Hello", msgid
    if in_reply_to:
        m["In-Reply-To"] = in_reply_to
    for k, v in headers.items():
        m[k.replace("_", "-")] = v
    m.set_content(body)
    return m.as_bytes()


def bounce_bytes(address, status="5.1.1", action="failed", msgid="<b1@mail.google.com>", quoted="<m1@me.org>"):
    return (f"From: Mail Delivery Subsystem <mailer-daemon@googlemail.com>\r\nTo: me@example.org\r\n"
            f"Subject: Delivery Status Notification (Failure)\r\nMessage-ID: {msgid}\r\n"
            f"Content-Type: multipart/report; report-type=delivery-status; boundary=\"BB\"\r\n\r\n"
            f"--BB\r\nContent-Type: text/plain\r\n\r\nAddress not found\r\n"
            f"--BB\r\nContent-Type: message/delivery-status\r\n\r\n"
            f"Final-Recipient: rfc822; {address}\r\nAction: {action}\r\nStatus: {status}\r\n\r\n"
            f"--BB\r\nContent-Type: message/rfc822\r\n\r\nMessage-ID: {quoted}\r\nSubject: Hello\r\n\r\nHi\r\n--BB--\r\n"
            ).encode()


class FakeIMAP:
    """Just enough of imaplib.IMAP4_SSL for check_replies: UID SEARCH / UID FETCH over a dict of messages."""
    box = {}            # uid -> raw bytes
    seen_fetches = []

    def __init__(self, host, timeout=None):
        pass

    def login(self, user, pw):
        return "OK", [b""]

    def select(self, name, readonly=False):
        assert readonly, "the inbox must be opened read-only"
        return "OK", [str(len(self.box)).encode()]

    def response(self, code):
        return code, [b"777"]

    def uid(self, cmd, *args):
        if cmd == "SEARCH":
            assert "UNSEEN" not in args, "read mail must be looked at too"
            return "OK", [" ".join(str(u) for u in sorted(self.box)).encode()]
        ids, what = args
        FakeIMAP.seen_fetches.append((ids, what))
        out = []
        for u in (int(x) for x in ids.split(",")):
            raw = self.box[u]
            data = raw.split(b"\r\n\r\n", 1)[0] + b"\r\n\r\n" if "HEADER" in what else raw
            if "HEADER" in what and b"\r\n\r\n" not in raw:
                data = raw.split(b"\n\n", 1)[0] + b"\n\n"
            out += [(f"{u} (UID {u} BODY[] {{{len(data)}}}".encode(), data), b")"]
        return "OK", out

    def logout(self):
        return "BYE", [b""]


class TestMailMergeAndPlaceholders(Base):
    def setUp(self):
        super().setUp()
        db.set_settings({"templates": [
            {"name": "site", "channel": "email", "body": "Subject: A site for {business_name}\n\nHi {first_name|there}, saw {business_name} in {city|your area}."},
            {"name": "wa", "channel": "whatsapp", "body": "Subject: ignored\nHi {name} team!"}]})

    def test_exact_mode_fills_the_template_and_never_calls_the_ai(self):
        lead_id = self.add(business_name="Smile Dental", email="a@smile.com", address="1 Main St, Dallas, TX 75201")
        with mock.patch.object(llm, "complete", side_effect=AssertionError("the AI must not be called")):
            out = outreach.draft_message(lead_id, exact=True)
        self.assertEqual(out["message"], "Subject: A site for Smile Dental\n\nHi there, saw Smile Dental in Dallas.")
        self.assertEqual(out["unfilled"], [])
        self.assertEqual(db.get_lead(lead_id)["template_used"], "site")

    def test_exact_mode_follows_the_setting_and_needs_no_ai_key(self):
        lead_id = self.add(business_name="Smile Dental", phone="+91 98765 43210")     # WhatsApp lead
        db.set_settings({"exact_templates": True})
        with mock.patch.dict(os.environ, {"GROQ_API_KEY": ""}), mock.patch.object(llm, "complete", side_effect=AssertionError):
            msg = outreach.run_draft(FakeCtx(), {})
        self.assertEqual(msg, "1 drafted")
        self.assertEqual(db.get_lead(lead_id)["message"], "Hi Smile Dental team!")       # the Subject line is dropped for a DM

    def test_a_blank_without_a_fallback_is_reported_and_blocks_the_send(self):
        db.set_settings({"templates": [{"name": "t", "channel": "email", "body": "Hi {first_name}, hello."}]})
        lead_id = self.add(business_name="Acme", email="a@acme.com")
        ctx = FakeCtx()
        self.assertEqual(outreach.run_draft(ctx, {"exact": True}), "1 drafted, 1 with a blank still to fill")
        self.assertTrue(any("{first_name}" in l for l in ctx.lines))
        with mock.patch.object(outreach, "_deliver") as deliver, self.assertRaises(outreach.SendError):
            outreach.send_lead_email(lead_id)
        deliver.assert_not_called()

    def test_no_template_for_the_channel_is_a_clear_failure_not_an_empty_draft(self):
        lead_id = self.add(business_name="Yelp Only", yelp_url="https://www.yelp.com/biz/yelp-only")
        ctx = FakeCtx()
        self.assertEqual(outreach.run_draft(ctx, {"exact": True}), "0 drafted, 1 failed")
        self.assertEqual(db.get_lead(lead_id)["message"], "")
        self.assertTrue(any("no Yelp template" in l for l in ctx.lines))

    def test_fallback_syntax(self):
        lead = {"business_name": "Acme", "contact_name": "", "city": ""}
        self.assertEqual(schema.fill_placeholders("Hi {first_name|there} / {{ Owner Name | friend }} / [first_name]", lead),
                         ("Hi there / friend / [first_name]", ["[first_name]"]))
        self.assertEqual(schema.fill_placeholders("Hi {first_name|there}", {"contact_name": "Dr. Asha Rao"})[0], "Hi Asha")
        self.assertEqual(schema.fill_placeholders("{first_name|}", lead), ("", []))       # an empty fallback: just leave it out

    def test_leads_with_no_way_to_reach_them_are_not_drafted(self):
        self.add(business_name="No Contact Co")
        with mock.patch.object(llm, "complete", return_value="hello") as ai:
            self.assertEqual(outreach.run_draft(FakeCtx(), {}), "0 drafted")
        ai.assert_not_called()


class TestCleaning(Base):
    def test_placeholder_cells_are_not_data(self):
        csv_text = ("Business Name,Phone,Email,Website,Facebook,Instagram,Notes\n"
                    "Real Biz,(214) 555-0101,N/A,none,No,-,call after 5\n"
                    "N/A,-,n/a,,,,\n")
        s = importer.import_csv_bytes(csv_text.encode())
        self.assertEqual((s["inserted"], s["skipped_empty"]), (1, 1))
        lead = db.list_leads({})[0][0]
        self.assertEqual((lead["email"], lead["website"], lead["facebook_url"], lead["instagram_url"]), ("", "", "", ""))
        self.assertEqual(lead["channel"], "phone")
        self.assertEqual(lead["notes"], "call after 5")

    def test_first_address_is_taken_from_a_cell_with_several(self):
        self.assertEqual(schema.norm_email("info@acme.com; sales@acme.com"), "info@acme.com")
        self.assertEqual(schema.norm_email("Jo Smith <Jo@Acme.com>"), "jo@acme.com")
        self.assertEqual(schema.norm_email("logo@2x.png, real@acme.com"), "real@acme.com")
        self.assertEqual(schema.norm_email("not an email"), "")

    def test_listing_and_booking_pages_are_not_a_website_of_their_own(self):
        for url in ("https://www.practo.com/bhopal/clinic/smile", "https://wa.me/919876543210", "https://www.justdial.com/x",
                    "https://www.tripadvisor.com/Restaurant-x", "https://booksy.com/en-us/123", "https://linktr.ee/acme"):
            self.assertEqual(schema.normalize_lead({"website": url})["has_website"], 0, url)
        self.assertEqual(schema.normalize_lead({"website": "https://smiledental.in"})["has_website"], 1)

    def test_default_country_code_reads_local_numbers(self):
        with mock.patch.dict(os.environ, {"DEFAULT_COUNTRY_CODE": "91"}):
            for raw in ("98765 43210", "098765 43210", "919876543210", "+91 98765 43210"):
                self.assertEqual(schema.norm_phone(raw), "+919876543210", raw)
            self.assertEqual(schema.derive_channel(schema.normalize_lead({"phone": "98765 43210"})), "whatsapp")
        self.assertEqual(schema.norm_phone("(214) 555-0100"), "+12145550100")           # the US default is unchanged
        self.assertEqual(schema.norm_phone("98765 43210"), "+19876543210")

    def test_csv_in_windows_encoding_semicolons_and_blank_top_lines(self):
        data = "\n\nName;Phone;City\nCafé Zoë;(214) 555-0101;Dallas\n".encode("cp1252")
        s = importer.import_csv_bytes(data)
        self.assertEqual(s["inserted"], 1)
        self.assertEqual(db.list_leads({})[0][0]["business_name"], "Café Zoë")

    def test_excel_workbook_imports_like_a_csv(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Clinic Name", "Phone", "Reviews"])
        ws.append(["Smile Dental Care", "(214) 555-0101", 12.0])
        ws.append([None, None, None])
        ws.append(["Bright Eyes Optical", "(214) 555-0102", 3])
        buf = io.BytesIO()
        wb.save(buf)
        s = importer.import_csv_bytes(buf.getvalue(), filename="friends list.xlsx")
        self.assertEqual((s["inserted"], s["missing_name"]), (2, 0))
        lead = next(l for l in db.list_leads({})[0] if l["business_name"] == "Smile Dental Care")
        self.assertEqual((lead["phone"], lead["review_count"]), ("+12145550101", 12))
        with self.assertRaises(ValueError):
            importer.import_csv_bytes(b"\xd0\xcf\x11\xe0old xls", filename="old.xls")

    def test_export_cannot_smuggle_a_formula_and_survives_a_round_trip(self):
        self.assertEqual(exports.safe_cell("=HYPERLINK(\"http://evil\")"), "'=HYPERLINK(\"http://evil\")")
        self.assertEqual(exports.safe_cell("@SUM(A1)"), "'@SUM(A1)")
        self.assertEqual(exports.safe_cell("-cmd|x"), "'-cmd|x")
        for fine in ("+12145550100", "-5", "Acme Plumbing", 7, ""):
            self.assertEqual(exports.safe_cell(fine), fine)
        s = importer.import_csv_bytes(b"Business Name,Phone\n'=Equals Plumbing,(214) 555-0101\n")
        self.assertEqual(db.list_leads({})[0][0]["business_name"], "=Equals Plumbing")

    def test_removing_a_do_not_contact_entry_frees_its_leads(self):
        a = self.add(business_name="A", email="stop@a.com", phone="(214) 555-0101")
        db.add_suppression("stop@a.com")
        self.assertEqual(db.get_lead(a)["do_not_contact"], 1)
        db.add_suppression("(214) 555-0101")
        db.remove_suppression("stop@a.com")
        self.assertEqual(db.get_lead(a)["do_not_contact"], 1)            # the phone entry still covers them
        db.remove_suppression("+12145550101")
        self.assertEqual(db.get_lead(a)["do_not_contact"], 0)

    def test_search_treats_percent_and_underscore_literally(self):
        self.add(business_name="100% Plumbing", phone="(214) 555-0101")
        self.add(business_name="Other Co", phone="(214) 555-0102")
        self.assertEqual(db.list_leads({"q": "%"})[1], 1)
        self.assertEqual(db.list_leads({"q": "_"})[1], 0)


class TestJudgeEnrichAI(Base):
    def test_a_reply_with_no_verdict_leaves_the_lead_unjudged(self):
        lead_id = self.add(business_name="Acme", phone="(214) 555-0101")
        with mock.patch.object(llm, "complete", return_value='{"confidence": "high", "reason": "looks fine"}'):
            msg = filtering.run_judge(FakeCtx(), {})
        self.assertIsNone(db.get_lead(lead_id)["fit"])
        self.assertIn("1 could not be judged", msg)
        with mock.patch.object(llm, "complete", return_value='{"fits": "no", "confidence": "high", "reason": "x"}'):
            filtering.run_judge(FakeCtx(), {})
        self.assertEqual(db.get_lead(lead_id)["fit"], 0)

    def test_enrich_survives_a_lead_deleted_mid_run(self):
        a = self.add(business_name="A", phone="(214) 555-0101")
        b = self.add(business_name="B", phone="(214) 555-0102")
        real = enrichment.enrich_lead

        def gone(lead, *a_, **k):
            db.delete_leads([b])
            return {"updates": {}, "notes": "", "searches": 0}
        with mock.patch.object(enrichment, "enrich_lead", side_effect=gone):
            msg = enrichment.run_enrich(FakeCtx(), {})
        self.assertIn("1 enriched", msg)
        self.assertIsNotNone(db.get_lead(a)["enriched_at"])
        self.assertIs(enrichment.enrich_lead, real)

    def test_rate_limit_waits_as_asked_and_does_not_sleep_after_the_last_try(self):
        class R:
            status_code, text, headers = 429, "slow down", {"Retry-After": "7"}
        with mock.patch.object(llm.requests, "post", return_value=R()), mock.patch.object(llm.time, "sleep") as sleep:
            with self.assertRaises(llm.LLMError) as cm:
                llm.complete("hi")
        self.assertEqual([c.args[0] for c in sleep.call_args_list], [7.0, 7.0])       # two waits for three tries
        self.assertIn("rate limit", str(cm.exception))


class TestSendingHoursAndToday(Base):
    def at(self, y, mo, d, h):
        return datetime(y, mo, d, h, 0, tzinfo=timezone.utc)

    def test_off_by_default(self):
        self.assertEqual(outreach.send_window(), (True, ""))

    def test_inside_and_outside_in_the_accounts_own_time_zone(self):
        db.set_settings({"send_window_enabled": True, "timezone": "America/Chicago", "send_days": [0, 1, 2, 3, 4],
                         "send_start_hour": 8, "send_end_hour": 18})
        s = db.get_settings()
        self.assertTrue(outreach.send_window(s, self.at(2026, 10, 13, 15))[0])        # Tue 10:00 in Chicago
        ok, why = outreach.send_window(s, self.at(2026, 10, 13, 8))                    # Tue 03:00 in Chicago
        self.assertFalse(ok)
        self.assertIn("03:00", why)
        ok, why = outreach.send_window(s, self.at(2026, 10, 11, 15))                   # a Sunday
        self.assertFalse(ok)
        self.assertIn("Sunday", why)

    def test_bulk_send_refuses_outside_hours_and_sends_nothing(self):
        lead_id = self.add(business_name="Acme", email="a@acme.com")
        db.update_lead(lead_id, {"message": "Subject: Hi\n\nHello"})
        db.set_settings({"send_window_enabled": True, "send_days": []})
        with mock.patch.object(outreach, "_deliver") as deliver, self.assertRaises(outreach.SendError):
            outreach.run_send_bulk(FakeCtx(), {})
        deliver.assert_not_called()

    def test_today_follows_the_accounts_calendar_day(self):
        lead_id = self.add(business_name="Acme", email="a@acme.com")
        now = datetime.now(timezone.utc)
        with db.get_conn() as conn:                     # an email sent 3 hours ago
            conn.execute("INSERT INTO events(lead_id, kind, detail, created_at) VALUES(?,?,?,?)",
                         (lead_id, "email_sent", "{}", (now - timedelta(hours=3)).isoformat(timespec="seconds")))
        # pick the zone where the local time is ~01:00, i.e. that email was "yesterday" there
        offset = (1 - now.hour) % 24
        offset = offset - 24 if offset > 14 else offset
        zone = f"Etc/GMT{-offset:+d}" if offset else "UTC"
        local_hour = (now.hour + offset) % 24
        self.assertEqual(local_hour, 1)
        db.set_settings({"timezone": zone})
        self.assertEqual(outreach.emails_sent_today(), 0)
        db.set_settings({"timezone": ""})
        self.assertEqual(outreach.emails_sent_today(), 1 if now.hour >= 3 else 0)


class TestDeliveryFailures(Base):
    def queue(self, n):
        ids = []
        for i in range(n):
            lead_id = self.add(business_name=f"Biz {i}", email=f"a{i}@biz{i}.com", review_count=n - i)
            db.update_lead(lead_id, {"message": "Subject: Hi\n\nHello"})
            ids.append(lead_id)
        db.set_settings({"email_send_delay_seconds": 0, "warmup_enabled": False})
        return ids

    def test_a_wrong_password_stops_the_run_and_marks_nobody_failed(self):
        ids = self.queue(4)
        with mock.patch.object(outreach, "_deliver", side_effect=outreach.DeliveryError("login refused")) as deliver:
            ctx = FakeCtx()
            msg = outreach.run_send_bulk(ctx, {})
        self.assertEqual(deliver.call_count, 2)                            # one, one more to be sure, then stop
        self.assertEqual(msg, "0 sent, 0 failed, 0 skipped")
        self.assertEqual({db.get_lead(i)["send_status"] for i in ids}, {"unsent"})
        self.assertTrue(any("Nothing was marked as failed" in l for l in ctx.lines))

    def test_a_refused_address_fails_that_lead_only(self):
        ids = self.queue(3)

        def deliver(m):
            if m["To"] == "a1@biz1.com":
                raise outreach.DeliveryError("550 no such user", recipient=True)
        with mock.patch.object(outreach, "_deliver", side_effect=deliver):
            msg = outreach.run_send_bulk(FakeCtx(), {})
        self.assertEqual(msg, "2 sent, 1 failed, 0 skipped")
        self.assertEqual([db.get_lead(i)["send_status"] for i in ids], ["sent", "failed", "sent"])

    def test_one_message_the_server_rejects_does_not_block_the_queue(self):
        ids = self.queue(3)
        calls = []

        def deliver(m):
            calls.append(m["To"])
            if len(calls) == 1:
                raise outreach.DeliveryError("554 message rejected")       # not address-specific, but only this one
        with mock.patch.object(outreach, "_deliver", side_effect=deliver):
            msg = outreach.run_send_bulk(FakeCtx(), {})
        self.assertEqual(msg, "2 sent, 1 failed, 0 skipped")
        self.assertEqual([db.get_lead(i)["send_status"] for i in ids], ["failed", "sent", "sent"])

    def test_smtp_errors_are_sorted_into_ours_and_theirs(self):
        import smtplib

        class S:
            def __init__(self, exc): self.exc = exc
            def __enter__(self): return self
            def __exit__(self, *a): return False
            def starttls(self): pass
            def login(self, u, p):
                if isinstance(self.exc, smtplib.SMTPAuthenticationError): raise self.exc
            def send_message(self, m): raise self.exc
        msg = EmailMessage()
        with mock.patch.object(outreach.netguard, "check_mail_host"):
            for exc, theirs in ((smtplib.SMTPRecipientsRefused({"x@y.com": (550, b"No such user")}), True),
                                (smtplib.SMTPAuthenticationError(535, b"bad"), False),
                                (smtplib.SMTPDataError(550, b"Daily user sending limit exceeded"), False),
                                (OSError("network unreachable"), False)):
                with mock.patch.object(outreach.smtplib, "SMTP", return_value=S(exc)):
                    with self.assertRaises(outreach.DeliveryError) as cm:
                        outreach._deliver(msg)
                self.assertEqual(cm.exception.recipient, theirs, exc)


class TestRepliesAndBounces(Base):
    def setUp(self):
        super().setUp()
        FakeIMAP.box, FakeIMAP.seen_fetches = {}, []
        self._p = [mock.patch.object(outreach.imaplib, "IMAP4_SSL", FakeIMAP),
                   mock.patch.object(outreach.netguard, "check_mail_host")]
        for p in self._p:
            p.start()

    def tearDown(self):
        for p in self._p:
            p.stop()
        super().tearDown()

    def test_a_reply_is_found_whether_or_not_it_was_already_read(self):
        lead_id = emailed(self)
        FakeIMAP.box = {5: reply_bytes("friend@elsewhere.com", in_reply_to=None, msgid="<x@e.com>"),
                        9: reply_bytes("Owner <owner@acme.com>")}
        c = outreach.check_replies()
        self.assertEqual((c["checked"], c["replied"]), (2, 1))
        self.assertEqual(db.get_lead(lead_id)["stage"], "replied")
        self.assertEqual(outreach.followups_due(), [])
        # only the message that could be ours was downloaded in full
        full = [f for f in FakeIMAP.seen_fetches if "HEADER" not in f[1]]
        self.assertEqual([f[0] for f in full], ["9"])

    def test_the_next_check_only_looks_at_new_mail(self):
        emailed(self)
        FakeIMAP.box = {5: reply_bytes("owner@acme.com")}
        outreach.check_replies()
        FakeIMAP.seen_fetches = []
        FakeIMAP.box[6] = reply_bytes("other@nowhere.com", in_reply_to=None, msgid="<n@n.com>")
        c = outreach.check_replies()
        self.assertEqual(c["checked"], 1)
        self.assertEqual(FakeIMAP.seen_fetches[0][0], "6")
        self.assertTrue(db.get_setting("last_reply_check_at"))

    def test_a_hard_bounce_marks_the_lead_failed_and_stops_follow_ups(self):
        lead_id = emailed(self)
        self.assertEqual(len(outreach.followups_due()), 1)
        FakeIMAP.box = {3: bounce_bytes("owner@acme.com")}
        c = outreach.check_replies()
        self.assertEqual((c["bounced"], c["replied"]), (1, 0))
        lead = db.get_lead(lead_id)
        self.assertEqual((lead["send_status"], lead["stage"]), ("failed", "contacted"))
        self.assertEqual(outreach.followups_due(), [])
        self.assertTrue(db.event_exists("bounced", "owner@acme.com"))

    def test_a_bounce_is_matched_by_the_quoted_message_when_no_address_is_given(self):
        lead_id = emailed(self)
        raw = bounce_bytes("someone-else@forwarded.com")           # the address in the report isn't the lead's
        self.assertEqual(outreach.process_inbound(raw), "bounced")
        self.assertEqual(db.get_lead(lead_id)["send_status"], "failed")

    def test_a_delay_notice_is_not_a_bounce(self):
        lead_id = emailed(self)
        raw = bounce_bytes("owner@acme.com", status="4.4.1", action="delayed")
        self.assertEqual(outreach.process_inbound(raw), "auto")
        self.assertEqual(db.get_lead(lead_id)["send_status"], "sent")

    def test_a_bounce_about_someone_elses_mail_is_ignored(self):
        emailed(self)
        self.assertIsNone(outreach.process_inbound(bounce_bytes("stranger@x.com", quoted="<zzz@other.org>")))

    def test_timer_checks_only_when_connected_switched_on_and_due(self):
        emailed(self)
        FakeIMAP.box = {5: reply_bytes("owner@acme.com")}
        with mock.patch.object(outreach, "imap_configured", return_value=False):
            self.assertIsNone(outreach.auto_check_replies())
        db.set_settings({"reply_check_minutes": 0})
        self.assertIsNone(outreach.auto_check_replies())
        db.set_settings({"reply_check_minutes": 30})
        self.assertEqual(outreach.auto_check_replies()["replied"], 1)
        self.assertIsNone(outreach.auto_check_replies())                   # just checked
        later = datetime.now(timezone.utc) + timedelta(minutes=31)
        self.assertIsNotNone(outreach.auto_check_replies(now=later))

    def test_replies_waiting_lists_who_to_answer_with_what_they_said(self):
        lead_id = emailed(self)
        outreach.process_inbound(reply_bytes("owner@acme.com", body="What would it cost?"))
        rows = db.replies_waiting()
        self.assertEqual([(r["id"], r["snippet"]) for r in rows], [(lead_id, "What would it cost?")])
        db.update_lead(lead_id, {"stage": "qualified"})
        self.assertEqual(db.replies_waiting(), [])


class TestFollowUps(Base):
    def setUp(self):
        super().setUp()
        db.set_settings({"email_send_delay_seconds": 0, "warmup_enabled": False, "max_followups": 2})

    def test_your_own_follow_up_text_is_used_with_blanks_filled_and_no_ai(self):
        lead_id = emailed(self)
        db.set_settings({"followup_template": "Hi {first_name|there}, just floating this back up for {business_name}."})
        sent = []
        with mock.patch.object(outreach, "_deliver", side_effect=sent.append), \
                mock.patch.object(llm, "complete", side_effect=AssertionError):
            outreach.send_followup(lead_id)
        self.assertIn("Hi there, just floating this back up for Acme Plumbing.", sent[0].get_content())
        self.assertEqual((sent[0]["Subject"], sent[0]["In-Reply-To"]), ("Re: Hello", "<m1@me.org>"))
        self.assertEqual(db.get_lead(lead_id)["follow_up_count"], 1)

    def test_send_all_due_sends_each_once_for_the_chosen_campaign(self):
        a = emailed(self, "A Co", "a@a.com", "one", "<a@me.org>")
        b = emailed(self, "B Co", "b@b.com", "two", "<b@me.org>")
        db.set_settings({"followup_template": "Any thoughts, {business_name}?"})
        sent = []
        with mock.patch.object(outreach, "_deliver", side_effect=sent.append), \
                mock.patch.object(outreach, "imap_configured", return_value=False):
            msg = outreach.run_followups(FakeCtx(), {"campaign": "one"})
        self.assertEqual(msg, "1 follow-ups sent, 0 failed, 0 skipped")
        self.assertEqual((db.get_lead(a)["follow_up_count"], db.get_lead(b)["follow_up_count"]), (1, 0))
        self.assertEqual(outreach.followups_due()[0]["id"], b)             # A was just nudged, so it's no longer due

    def test_no_follow_up_text_and_no_ai_stops_with_a_reason(self):
        emailed(self)
        ctx = FakeCtx()
        with mock.patch.dict(os.environ, {"GROQ_API_KEY": ""}), mock.patch.object(outreach, "_deliver") as deliver, \
                mock.patch.object(outreach, "imap_configured", return_value=False):
            outreach.run_followups(ctx, {})
        deliver.assert_not_called()
        self.assertTrue(any("no follow-up message to send" in l for l in ctx.lines))

    def test_leads_messaged_by_hand_come_due_and_are_recorded_as_followed_up(self):
        lead_id = self.add(business_name="Smile Dental", phone="+91 98765 43210")
        db.update_lead(lead_id, {"message": "Hi Smile Dental team!"})
        outreach.mark_sent(lead_id)
        self.assertEqual(outreach.hand_followups_due(), [])                # sent a moment ago
        db.update_lead(lead_id, {"last_sent_at": OLD})
        self.assertEqual([l["id"] for l in outreach.hand_followups_due()], [lead_id])
        self.assertEqual(outreach.followups_due(), [])                     # it's not an email follow-up
        db.set_settings({"followup_template": "Subject: x\nHi {name}, any thoughts?"})
        self.assertEqual(outreach.followup_text(db.get_lead(lead_id)), "Hi Smile Dental, any thoughts?")
        outreach.mark_sent(lead_id, followup=True)
        lead = db.get_lead(lead_id)
        self.assertEqual((lead["follow_up_count"], lead["stage"], lead["send_status"]), (1, "contacted", "sent"))
        self.assertEqual(outreach.hand_followups_due(), [])
        self.assertTrue(any(e["kind"] == "manual_followup" for e in db.list_events(lead_id)))


class TestPipeline(Base):
    def test_enrich_then_judge_then_draft_only_the_ones_that_fit(self):
        good = self.add(business_name="Good Co", email="a@good.com", campaign="c")
        bad = self.add(business_name="Bad Co", email="b@bad.com", campaign="c")
        other = self.add(business_name="Elsewhere", email="c@else.com", campaign="other")

        def ai(prompt, task="draft", **k):
            if task == "filter":
                return json.dumps({"fits": "Good Co" in prompt, "confidence": "high", "reason": "r"})
            return "Subject: Hi\n\nHello Good Co"
        ctx = FakeCtx()
        with mock.patch.object(llm, "complete", side_effect=ai), \
                mock.patch.object(enrichment, "enrich_lead", return_value={"updates": {}, "notes": "", "searches": 0}):
            msg = pipeline.run_pipeline(ctx, {"campaign": "c"})
        self.assertIn("Enrich: 2 enriched", msg)
        self.assertIn("Judge: 1 fit, 1 rejected", msg)
        self.assertIn("Draft: 1 drafted", msg)
        self.assertTrue(db.get_lead(good)["message"])
        self.assertEqual(db.get_lead(bad)["message"], "")
        self.assertIsNone(db.get_lead(other)["enriched_at"])               # the other campaign wasn't touched
        self.assertTrue(any(l.startswith("[judge]") for l in ctx.lines))
        with mock.patch.object(llm, "complete", side_effect=AssertionError("nothing left to do")), \
                mock.patch.object(enrichment, "enrich_lead", side_effect=AssertionError):
            self.assertIn("0 enriched", pipeline.run_pipeline(FakeCtx(), {"campaign": "c"}))

    def test_without_an_ai_key_it_still_enriches_and_says_what_was_skipped(self):
        self.add(business_name="Good Co", email="a@good.com")
        with mock.patch.dict(os.environ, {"GROQ_API_KEY": ""}), \
                mock.patch.object(enrichment, "enrich_lead", return_value={"updates": {}, "notes": "", "searches": 0}):
            msg = pipeline.run_pipeline(FakeCtx(), {})
        self.assertEqual(msg, "Enrich: 1 enriched, 0 new fields found, 0 searches used | Judge: skipped (no AI key) | Draft: skipped (no AI key)")


class TestApiAdditions(ApiBase):
    def test_schedules_keep_the_no_website_choice_and_get_the_same_checks(self):
        r = self.client.post("/api/schedules", json={"platform": "google_maps", "query": "AC repair; furnace repair",
                                                     "location": "Garland, TX", "max_results": 9999,
                                                     "only_no_website": True, "interval_hours": 12})
        self.assertEqual(r.status_code, 200, r.get_json())
        s = db.list_schedules()[0]
        self.assertEqual((s["only_no_website"], s["max_results"], s["query"]), (1, 200, "AC repair; furnace repair"))
        launched = []
        jobs.run_scheduler_once(lambda *a: launched.append(a))
        self.assertEqual(launched[0][-1], True)                             # the filter reaches the scrape
        too_many = ";".join(f"t{i}" for i in range(41))
        self.assertEqual(self.client.post("/api/schedules", json={"platform": "google_maps", "query": too_many}).status_code, 400)

    def test_a_full_queue_leaves_the_schedule_due_and_one_bad_user_does_not_stop_the_next(self):
        db.create_schedule("google_maps", "plumbers", "Dallas, TX", 10, "c", 24)
        self.signup("second@example.com")
        with self.accounts_ctx(2):
            db.create_schedule("google_maps", "roofers", "Dallas, TX", 10, "c", 24)
        launched = []

        def launch(platform, query, *a):
            if query == "plumbers":
                raise jobs.JobBusy("queue full")
            launched.append(query)
        jobs.run_scheduler_once(launch)
        self.assertEqual(launched, ["roofers"])
        self.assertEqual(len(db.due_schedules()), 1)                        # the owner's is still due

    def accounts_ctx(self, uid):
        import userctx
        return userctx.use(self.accounts.get_user(uid), self.accounts.db_path_for(uid))

    def test_job_list_sends_log_tails_unless_asked_for_the_whole_log(self):
        j = db.create_job("scrape", {})
        db.update_job(j, log="\n".join(f"line {i}" for i in range(200)))
        full = self.client.get("/api/jobs").get_json()[0]
        self.assertEqual(len(full["log"].split("\n")), 200)                 # old callers still get everything
        light = self.client.get("/api/jobs?logs=").get_json()[0]
        self.assertEqual(light["log"].split("\n"), [f"line {i}" for i in range(188, 200)])
        self.assertTrue(light["log_more"])
        opened = self.client.get(f"/api/jobs?logs={j}").get_json()[0]
        self.assertEqual((len(opened["log"].split("\n")), opened["log_more"]), (200, False))

    def test_the_browser_check_is_asked_once_not_on_every_page_load(self):
        import scraper
        with mock.patch.object(scraper, "browser_available", return_value=True) as probe:
            for _ in range(5):
                self.assertTrue(self.client.get("/api/config-status").get_json()["playwright"])
        self.assertEqual(probe.call_count, 1)

    def test_unexpected_errors_come_back_as_json_with_a_message(self):
        with mock.patch.object(db, "stats", side_effect=RuntimeError("boom")), mock.patch("traceback.print_exc"):
            r = self.client.get("/api/stats")
        self.assertEqual(r.status_code, 500)
        self.assertIn("Something went wrong", r.get_json()["error"])
        self.assertEqual(self.client.delete("/api/stats").status_code, 405)
        self.assertIn("error", self.client.delete("/api/stats").get_json())

    def test_lead_list_filters_and_limits(self):
        a = self.add(business_name="Emailable", email="a@a.com")
        b = self.add(business_name="Phone Only", phone="(214) 555-0101")
        db.update_lead(a, {"message": "hi"})
        get = lambda q: [l["business_name"] for l in self.client.get("/api/leads?" + q).get_json()["leads"]]
        self.assertEqual(get("has=email"), ["Emailable"])
        self.assertEqual(get("has=no_email"), ["Phone Only"])
        self.assertEqual(get("sent=drafted"), ["Emailable"])
        self.assertEqual(get("sent=none"), ["Phone Only"])
        self.assertEqual(get("hand=1"), ["Phone Only"])
        self.assertEqual(len(get("limit=-5")), 1)                           # a nonsense limit can't mean "everything"
        self.assertEqual(self.client.get("/api/leads?limit=abc").status_code, 200)

    def test_draft_route_needs_an_ai_key_only_when_the_ai_is_writing(self):
        self.add(business_name="Acme", email="a@acme.com")
        db.set_settings({"templates": [{"name": "t", "channel": "email", "body": "Hi {business_name}"}]})
        with mock.patch.dict(os.environ, {"GROQ_API_KEY": ""}):
            self.assertEqual(self.client.post("/api/draft", json={}).status_code, 400)
            r = self.client.post("/api/draft", json={"exact": True})
            self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(db.list_leads({})[0][0]["message"], "Hi Acme")

    def test_settings_validation_for_hours_zone_and_days(self):
        put = lambda b: self.client.put("/api/settings", json=b)
        self.assertEqual(put({"timezone": "Mars/Olympus"}).status_code, 400)
        self.assertEqual(put({"send_start_hour": 18, "send_end_hour": 9}).status_code, 400)
        self.assertEqual(put({"email_daily_limit": "lots"}).status_code, 400)
        r = put({"timezone": "Asia/Kolkata", "send_window_enabled": True, "send_days": [0, 2, 9, "x"],
                 "send_start_hour": 9, "send_end_hour": 17, "exact_templates": True,
                 "followup_template": " Any thoughts? ", "reply_check_minutes": 15})
        self.assertEqual(r.status_code, 200, r.get_json())
        s = r.get_json()
        self.assertEqual((s["timezone"], s["send_days"], s["followup_template"], s["reply_check_minutes"]),
                         ("Asia/Kolkata", [0, 2], "Any thoughts?", 15))

    def test_changing_the_country_code_re_reads_numbers_it_could_not_understand(self):
        lead_id = self.add(business_name="Smile Dental", phone="98765 4321")       # 9 digits: no good as a US number
        self.assertEqual(db.get_lead(lead_id)["phone"], "")
        r = self.client.put("/api/connections", json={"set": {"DEFAULT_COUNTRY_CODE": "91"}})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(r.get_json()["phones_fixed"], 1)
        lead = db.get_lead(lead_id)
        self.assertEqual((lead["phone"], lead["channel"]), ("+91987654321", "whatsapp"))
        self.assertEqual(self.client.put("/api/connections", json={"set": {"DEFAULT_COUNTRY_CODE": "abc"}}).status_code, 400)

    def test_scrape_can_hand_over_to_the_pipeline(self):
        import scraper
        with mock.patch.object(scraper, "run_scrape", return_value="3 scraped -> 3 new, 0 already known"), \
                mock.patch.object(scraper, "browser_available", return_value=True), \
                mock.patch.object(pipeline, "run_pipeline", return_value="Enrich: 3 enriched") as pipe:
            r = self.client.post("/api/scrape", json={"platform": "google_maps", "query": "plumbers",
                                                      "campaign": "c", "then_pipeline": True})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(pipe.call_args[0][1]["campaign"], "c")
        kinds = {j["kind"]: j for j in db.list_jobs()}
        self.assertIn("now enriching", kinds["scrape"]["summary"])
        self.assertEqual(kinds["pipeline"]["summary"], "Enrich: 3 enriched")


if __name__ == "__main__":
    unittest.main()


class TestUpgradingAnOlderDatabase(Base):
    def test_schedules_from_before_the_no_website_column_still_work(self):
        import sqlite3
        import tempfile
        path = tempfile.NamedTemporaryFile(suffix=".db", delete=False).name
        c = sqlite3.connect(path)
        c.executescript("""CREATE TABLE schedules (id INTEGER PRIMARY KEY AUTOINCREMENT, platform TEXT, query TEXT,
            location TEXT DEFAULT '', max_results INTEGER DEFAULT 30, campaign TEXT DEFAULT '', interval_hours INTEGER DEFAULT 24,
            last_run_at TEXT, next_run_at TEXT, active INTEGER DEFAULT 1);
            INSERT INTO schedules(platform, query, next_run_at) VALUES('google_maps', 'plumbers', '2020-01-01T00:00:00+00:00');""")
        c.commit()
        c.close()
        db.set_path(path)
        try:
            db.init_db()
            self.assertEqual(db.due_schedules()[0]["only_no_website"], 0)
            db.create_schedule("google_maps", "roofers", "", 30, "", 24, only_no_website=True)
            self.assertEqual(db.list_schedules()[0]["only_no_website"], 1)
        finally:
            db.set_path(self._tmp.name)
            os.unlink(path)

    def test_old_runs_are_pruned_but_recent_ones_kept(self):
        for _ in range(320):
            db.create_job("enrich", {})
        db.init_db()
        with db.get_conn() as conn:
            n, lo = conn.execute("SELECT COUNT(*), MIN(id) FROM jobs").fetchone()
        self.assertEqual((n, lo), (300, 21))
