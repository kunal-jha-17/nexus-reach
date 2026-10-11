"""
importer.py -- get leads into the hub from any CSV.

Headers are matched automatically: first by exact name (business_name, Business
Name, full_address, location_link ...), then by simple keyword rules. Columns
that can't be mapped are appended to the lead's notes (short values only) and
reported back, so nothing is silently lost.
"""
import csv
import io
import re

import db
import schema

MAX_IMPORT_ROWS = 20000

# normalised header -> canonical field
EXACT = {
    # business
    "business_name": "business_name", "business": "business_name", "name": "business_name",
    "company": "business_name", "company_name": "business_name", "title": "business_name",
    "organization": "business_name", "organisation": "business_name",
    # person
    "owner_name": "contact_name", "owner": "contact_name", "contact_name": "contact_name",
    "contact": "contact_name", "contact_person": "contact_name", "manager": "contact_name",
    "full_name": "contact_name",
    # category / trade
    "trade": "trade", "category": "category_raw", "category_raw": "category_raw",
    "type": "category_raw", "industry": "category_raw", "main_category": "category_raw",
    "subtypes": "category_raw",
    # location
    "address": "address", "full_address": "address", "street_address": "address",
    "city": "city", "state": "state", "us_state": "state",
    # reviews
    "reviews": "review_count", "review_count": "review_count", "reviews_count": "review_count",
    "user_ratings_total": "review_count", "number_of_reviews": "review_count",
    # contact
    "phone": "phone", "phone_1": "phone", "phone_number": "phone", "telephone": "phone",
    "tel": "phone", "mobile": "phone", "contact_number": "phone", "whatsapp": "phone",
    "email": "email", "email_1": "email", "e_mail": "email", "email_address": "email",
    "website": "website", "site": "website", "url": "website", "website_url": "website",
    "domain": "website", "web": "website",
    # social / listings
    "instagram": "instagram_url", "instagram_url": "instagram_url", "ig": "instagram_url",
    "ig_handle": "instagram_url", "insta": "instagram_url",
    "facebook": "facebook_url", "facebook_url": "facebook_url", "fb": "facebook_url",
    "fb_page": "facebook_url",
    "linkedin": "linkedin_url", "linkedin_url": "linkedin_url",
    "yelp_url": "yelp_url", "yelp": "yelp_url", "yelp_link": "yelp_url",
    "google_business_url": "google_maps_url", "google_maps_url": "google_maps_url",
    "location_link": "google_maps_url", "maps_url": "google_maps_url",
    "google_url": "google_maps_url", "place_url": "google_maps_url",
    # misc
    "notes": "notes", "note": "notes", "description": "notes",
    "source": "source_label", "campaign": "campaign",
}

# fallbacks for messy headers (substring rules, checked in this order)
KEYWORDS = [
    ("email", "email"), ("phone", "phone"), ("instagram", "instagram_url"),
    ("facebook", "facebook_url"), ("linkedin", "linkedin_url"), ("yelp", "yelp_url"),
    ("website", "website"), ("address", "address"), ("categor", "category_raw"),
    ("review", "review_count"), ("business name", "business_name"),
]

# Pre-written outreach copy that may already sit in a spreadsheet. Imported as the
# lead's draft message when it matches the lead's channel.
MESSAGE_HEADERS = {
    "sms_whatsapp": "phone", "sms": "phone", "whatsapp_message": "phone", "text_message": "phone",
    "facebook_dm": "facebook", "facebook_message": "facebook",
    "instagram_dm": "instagram", "instagram_message": "instagram",
    "email_body": "email", "email_message": "email",
    "yelp_message": "yelp",
}
SUBJECT_HEADERS = ("email_subject", "subject")
# Words that mean a header holds message text, never contact data.
_TEXTY = ("subject", "body", "message", "dm", "template", "draft", "reply")

# Columns we know are handled/derived elsewhere -- don't dump into notes.
IGNORED = {"has_website", "id", "quality_score", "score", "status", "stage", "fit_confidence",
           "fit_reason", "confidence_notes", "score_reasons"}


_HAS_CONTACT = ("phone", "phone_raw", "email", "website", "instagram_url", "facebook_url",
                "yelp_url", "google_maps_url")


def _worth_keeping(norm):
    """A row is a lead if, once cleaned ("N/A" and "-" are nothing), it has a name or any way to reach it."""
    return bool(norm["business_name"] or any(norm.get(k) for k in _HAS_CONTACT))


def _cell(row, h):
    """One CSV cell as text. Undoes the leading apostrophe our own export puts in front of
    values a spreadsheet would otherwise run as a formula ('=..., '@...)."""
    v = (row.get(h) or "").strip()
    if len(v) > 1 and v[0] == "'" and v[1] in "=@+-":
        v = v[1:]
    return "" if v.lower() == "nan" else v


def _norm_header(h):
    return re.sub(r"[^a-z0-9]+", "_", (h or "").strip().lower()).strip("_")


def map_headers(headers):
    """Return ({header: canonical_field}, [unmapped headers])."""
    mapping, used_targets, unmapped = {}, set(), []
    # A file with an explicit business column (Clinic Name, Company, Shop ...) AND a bare
    # "Name" / "Title" column: the bare one is the person / their job title, not the business.
    explicit_business = any(
        k not in ("name", "title") and (EXACT.get(k) == "business_name"
                                        or schema.name_header_target(k) == "business_name")
        for k in map(_norm_header, headers))
    for h in headers:
        key = _norm_header(h)
        if key in MESSAGE_HEADERS or key in SUBJECT_HEADERS:
            continue  # handled separately as draft messages
        if key in IGNORED:
            unmapped.append(h)
            continue
        target = EXACT.get(key)
        if explicit_business and key in ("name", "title"):
            target = "contact_name" if key == "name" else None
        if target is None and not any(w in key.split("_") for w in _TEXTY):
            # any "<something> name" column: Clinic Name, Name of Practice, Shop, Owner Name ...
            target = schema.name_header_target(key)
        if target is None and not any(w in key.split("_") for w in _TEXTY):
            for kw, tgt in KEYWORDS:
                if kw.replace(" ", "_") in key or kw in key:
                    target = tgt
                    break
        # first column wins for each target, except notes which can stack
        if target and (target not in used_targets or target == "notes"):
            mapping[h] = target
            used_targets.add(target)
        else:
            unmapped.append(h)
    return mapping, unmapped


def _decode(data):
    """Spreadsheet exports aren't always UTF-8: Excel on Windows writes Windows-1252 ("café"
    arrives as a broken character otherwise), and "Unicode text" is UTF-16."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _delimiter(text):
    """Comma, semicolon (Excel in much of Europe), tab or bar -- whichever the header row uses most."""
    first = next((ln for ln in text.splitlines() if ln.strip()), "")
    counts = {d: first.count(d) for d in (",", ";", "\t", "|")}
    best = max(counts, key=counts.get)
    return best if counts[best] else ","


def rows_from_csv_bytes(data):
    text = _decode(data)
    lines = text.splitlines(keepends=True)
    while lines and not lines[0].strip(" ,;\t|\r\n"):       # blank lines above the header row
        lines.pop(0)
    text = "".join(lines)
    reader = csv.DictReader(io.StringIO(text), delimiter=_delimiter(text))
    headers = [h for h in (reader.fieldnames or []) if h is not None and str(h).strip()]
    return headers, list(reader)


def rows_from_xlsx_bytes(data):
    """The first sheet of an Excel workbook, as (headers, rows) -- same shape as a CSV."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("This server can't read Excel files yet (the 'openpyxl' package is missing). "
                         "In Excel choose File > Save As > CSV and import that instead.")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 -- not a real workbook
        raise ValueError("That file couldn't be opened as an Excel workbook. Save it as CSV and try again.")
    try:
        ws = wb.worksheets[0]
        headers, rows = None, []
        for values in ws.iter_rows(values_only=True):
            cells = ["" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip()
                     for v in values]
            if not any(cells):
                continue
            if headers is None:
                headers = [c or f"column {i + 1}" for i, c in enumerate(cells)]
                continue
            rows.append(dict(zip(headers, cells)))
            if len(rows) > MAX_IMPORT_ROWS:
                break
        return headers or [], rows
    finally:
        wb.close()


def import_rows(headers, rows, campaign="", source="import", query=""):
    """Import parsed CSV rows. Returns a summary dict."""
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValueError(f"That file has {len(rows):,} rows; the limit is {MAX_IMPORT_ROWS:,} per import. Split it into smaller files.")
    mapping, unmapped = map_headers(headers)
    unmapped = [h for h in unmapped if _norm_header(h) not in MESSAGE_HEADERS
                and _norm_header(h) not in SUBJECT_HEADERS]
    msg_cols = {h: MESSAGE_HEADERS[_norm_header(h)] for h in headers if _norm_header(h) in MESSAGE_HEADERS}
    subj_col = next((h for h in headers if _norm_header(h) in SUBJECT_HEADERS), None)
    summary = {"inserted": 0, "merged": 0, "skipped_empty": 0, "skipped_suppressed": 0,
               "column_mapping": mapping,
               "unmapped_columns": [h for h in unmapped if _norm_header(h) not in IGNORED],
               "missing_name": 0, "ids": []}
    for start in range(0, len(rows), 200):
        with db.batch():
            for row in rows[start:start + 200]:
                lead = {}
                row_campaign, row_source = campaign, source
                extra_notes = []
                for h, target in mapping.items():
                    v = _cell(row, h)
                    if not v:
                        continue
                    if target == "campaign":
                        row_campaign = row_campaign or v
                    elif target == "source_label":
                        row_source = v
                    elif target == "notes":
                        extra_notes.append(v)
                    elif target in lead and lead[target]:
                        continue
                    else:
                        lead[target] = v
                for h in unmapped:
                    v = _cell(row, h)
                    if v and len(v) <= 200 and _norm_header(h) not in IGNORED and schema.clean_cell(h, v):
                        extra_notes.append(f"{h}: {v}")
                lead["notes"] = "; ".join(extra_notes)

                # Nothing to go on -> skip
                norm = schema.normalize_lead(lead)
                if not _worth_keeping(norm):
                    summary["skipped_empty"] += 1
                    continue
                if db.is_suppressed(norm):
                    summary["skipped_suppressed"] += 1
                    continue
                try:
                    lead_id, is_new = db.upsert_lead(lead, campaign=row_campaign, source=row_source, query=query)
                except db.LimitError as e:
                    summary["stopped_at_limit"] = str(e)
                    break
                summary["inserted" if is_new else "merged"] += 1
                summary["ids"].append(lead_id)
                if not norm["business_name"]:
                    summary["missing_name"] += 1

                # carry over any ready-made message that matches this lead's channel
                if msg_cols:
                    saved = db.get_lead(lead_id)
                    if saved and not saved["message"]:
                        want = "phone" if saved["channel"] == "whatsapp" else saved["channel"]
                        for h, ch in msg_cols.items():
                            text = (row.get(h) or "").strip()
                            if text and ch == want:
                                if ch == "email" and subj_col and (row.get(subj_col) or "").strip():
                                    text = f"Subject: {row[subj_col].strip()}\n\n{text}"
                                db.update_lead(lead_id, {"message": text})
                                break
        if summary.get("stopped_at_limit"):
            break
    return summary


def import_leads(leads, campaign="", source="", query=""):
    """Import already-structured leads (list of canonical-field dicts), e.g.
    straight from a scraper. Same dedupe + do-not-contact rules as CSV import."""
    summary = {"inserted": 0, "merged": 0, "skipped_empty": 0, "skipped_suppressed": 0, "ids": []}
    for start in range(0, len(leads), 200):
        with db.batch():
            for lead in leads[start:start + 200]:
                norm = schema.normalize_lead(lead)
                if not _worth_keeping(norm):
                    summary["skipped_empty"] += 1
                    continue
                if db.is_suppressed(norm):
                    summary["skipped_suppressed"] += 1
                    continue
                try:
                    lead_id, is_new = db.upsert_lead(lead, campaign=campaign, source=source, query=query)
                except db.LimitError as e:
                    summary["stopped_at_limit"] = str(e)
                    break
                summary["inserted" if is_new else "merged"] += 1
                summary["ids"].append(lead_id)
        if summary.get("stopped_at_limit"):
            break
    return summary


def import_csv_bytes(data, campaign="", source="import", filename=""):
    """Import a CSV -- or, going by the file name / its first bytes, an Excel .xlsx workbook."""
    if filename.lower().endswith((".xlsx", ".xlsm")) or data[:4] == b"PK\x03\x04":
        headers, rows = rows_from_xlsx_bytes(data)
    elif filename.lower().endswith(".xls") or data[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError("Old-style .xls files can't be read. In Excel choose File > Save As > "
                         "Excel Workbook (.xlsx) or CSV, then import that.")
    else:
        headers, rows = rows_from_csv_bytes(data)
    if not headers:
        raise ValueError("The file has no header row")
    return import_rows(headers, rows, campaign=campaign, source=source)
